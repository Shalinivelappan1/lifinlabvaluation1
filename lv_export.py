"""Exports: Excel workbook of the full valuation, JSON save/load of all inputs."""
from __future__ import annotations

import io
import json
from dataclasses import asdict

import pandas as pd

from lv_engine import Company, Assumptions
from lv_wacc import WaccInputs


def to_json(c: Company, a: Assumptions, narrative: dict, extras: dict) -> bytes:
    payload = {"version": 1, "company": asdict(c), "assumptions": asdict(a),
               "narrative": narrative, "extras": extras}
    return json.dumps(payload, indent=2, default=float).encode()


def from_json(raw: bytes):
    d = json.loads(raw)
    c = Company(**d["company"])
    ad = d["assumptions"]
    w = WaccInputs(**ad.pop("wacc"))
    a = Assumptions(**ad, wacc=w)
    return c, a, d.get("narrative", {}), d.get("extras", {})


def to_excel(sheets: dict) -> bytes:
    """sheets = {sheet name: DataFrame}. Percent-like columns are formatted automatically."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets.items():
            if df is None:
                continue
            df.to_excel(xw, sheet_name=name[:31])
            ws = xw.sheets[name[:31]]
            ws.column_dimensions["A"].width = 34
            for col in ws.iter_cols(min_col=2, max_col=ws.max_column):
                ws.column_dimensions[col[0].column_letter].width = 16
                header = str(col[0].value or "")
                for cell in col[1:]:
                    if isinstance(cell.value, (int, float)):
                        pct = any(k in header for k in ("growth", "margin", "rate", "WACC", "ROIC", "%",
                                                        "Probability", "ROE"))
                        cell.number_format = "0.0%" if pct and abs(cell.value) < 5 else "#,##0.00"
            from openpyxl.styles import Font, PatternFill
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="15233F")
    return buf.getvalue()
